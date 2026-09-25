import sys
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if len(sys.argv) < 3:
    print("Usage: python json_to_excel.py <input_json> <output_xlsx>")
    sys.exit(1)

json_path = sys.argv[1]
output_path = sys.argv[2]

try:
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    wb = openpyxl.Workbook()
    wb.remove(wb.active) # remove default sheet

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for table_name, rows in data.items():
        sheet_title = table_name[:31] # Excel sheet limit
        sheet = wb.create_sheet(title=sheet_title)
        if not rows:
            sheet.append(["id"])
            continue

        headers = list(rows[0].keys())
        sheet.append(headers)

        # Header styling
        for col_num in range(1, len(headers) + 1):
            cell = sheet.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data rows
        for row_data in rows:
            row_values = []
            for h in headers:
                val = row_data.get(h)
                if val is None:
                    row_values.append("")
                else:
                    row_values.append(val)
            sheet.append(row_values)

        # Styling & column auto-width
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = thin_border

        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    wb.save(output_path)
    print("SUCCESS")
except Exception as e:
    print(f"ERROR: {str(e)}")
    sys.exit(1)
