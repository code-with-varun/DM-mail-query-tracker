import sys
import json
import openpyxl
import datetime

if len(sys.argv) < 2:
    print("Usage: python excel_to_json.py <input_xlsx>")
    sys.exit(1)

excel_path = sys.argv[1]

try:
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    result = {}

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            continue

        headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
        sheet_data = []

        for row in rows[1:]:
            if not any(v is not None and str(v).strip() != '' for v in row):
                continue
            row_dict = {}
            for i, val in enumerate(row):
                if i >= len(headers):
                    continue
                h_name = headers[i]
                if isinstance(val, (datetime.datetime, datetime.date)):
                    if isinstance(val, datetime.datetime):
                        val = val.strftime('%Y-%m-%d %H:%M:%S')
                    else:
                        val = val.strftime('%Y-%m-%d')
                elif val is None or str(val).strip().upper() in ['NULL', 'NONE']:
                    val = None
                row_dict[h_name] = val
            sheet_data.append(row_dict)

        if sheet_data:
            result[sheet_name] = sheet_data

    print(json.dumps(result, ensure_ascii=False))
except Exception as e:
    sys.stderr.write(f"ERROR: {str(e)}\n")
    sys.exit(1)
